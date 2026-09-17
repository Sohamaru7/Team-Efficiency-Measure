"""API-level tests for the Phase 8 endpoints: POST /api/import/preview, POST /api/import/commit,
GET /api/import/history(/{id}), GET /api/export/tasks, GET /api/export/dashboard. Uploads real
CSV bytes through the FastAPI test client (multipart/form-data), exercising the full stack.
"""

import csv
import io

from openpyxl import load_workbook

from app.models.enums import ProjectStatus, UserRole
from app.models.project import Project
from app.models.task import Task
from app.models.user import User


def _seed(db_session):
    manager = User(name="Manager API Import", email="mgr-api-import@example.com", role=UserRole.MANAGER, department="Engineering")
    alice = User(name="Alice API Import", email="alice-api-import@example.com", role=UserRole.EMPLOYEE, department="Engineering")
    db_session.add_all([manager, alice])
    db_session.commit()
    project = Project(name="API Import Project", manager_id=manager.id, status=ProjectStatus.ACTIVE)
    db_session.add(project)
    db_session.commit()
    return {"manager": manager, "alice": alice, "project": project}


def _csv_bytes(project_name: str) -> bytes:
    return (
        "Employee,Task,Project,Priority,Estimated Hours,Start Date,Deadline,Status\n"
        f"Alice API Import,Design homepage,{project_name},high,5,2026-08-01,2026-08-10,in_progress\n"
    ).encode("utf-8")


def test_preview_endpoint_returns_result_without_writing(client, db_session):
    seed = _seed(db_session)
    resp = client.post(
        "/api/import/preview",
        files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["dry_run"] is True
    assert body["rows_accepted"] == 1
    assert body["rows"][0]["status"] == "accepted"
    assert db_session.query(Task).count() == 0


def test_commit_endpoint_creates_task_and_history(client, db_session):
    seed = _seed(db_session)
    resp = client.post(
        "/api/import/commit",
        files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")},
        data={"imported_by": str(seed["manager"].id)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["dry_run"] is False
    assert body["rows_accepted"] == 1
    assert body["history_id"] is not None
    assert db_session.query(Task).count() == 1


def test_commit_endpoint_attributes_import_to_the_authenticated_caller(client, db_session):
    """`imported_by` is derived from the real authenticated user (production-readiness pass) —
    there is no longer a client-supplied `imported_by` field to test rejecting; this asserts
    the history row is attributed to whoever the `client` fixture is authenticated as.
    """
    from app.models.import_history import ImportHistory
    from tests.conftest import TEST_ADMIN_EMAIL
    from app.models.user import User

    seed = _seed(db_session)
    resp = client.post("/api/import/commit", files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")})
    assert resp.status_code == 200
    history_id = resp.json()["history_id"]

    admin_id = db_session.query(User).filter_by(email=TEST_ADMIN_EMAIL).one().id
    record = db_session.get(ImportHistory, history_id)
    assert record.imported_by == admin_id


def test_commit_endpoint_reports_file_level_error(client, db_session):
    resp = client.post(
        "/api/import/commit",
        files={"file": ("tasks.csv", b"Employee,Priority\nAlice,high\n", "text/csv")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["file_error"] is not None
    assert body["history_id"] is not None


def test_commit_endpoint_rejects_unsupported_file_type_gracefully(client, db_session):
    resp = client.post(
        "/api/import/commit",
        files={"file": ("tasks.txt", b"not a real file", "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json()["file_error"] is not None


def test_import_history_list_and_get(client, db_session):
    seed = _seed(db_session)
    commit_resp = client.post(
        "/api/import/commit",
        files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")},
    )
    history_id = commit_resp.json()["history_id"]

    list_resp = client.get("/api/import/history")
    assert list_resp.status_code == 200
    assert any(h["id"] == history_id for h in list_resp.json())

    get_resp = client.get(f"/api/import/history/{history_id}")
    assert get_resp.status_code == 200
    assert get_resp.json()["filename"] == "tasks.csv"


def test_import_history_404_for_missing_id(client):
    resp = client.get("/api/import/history/999999")
    assert resp.status_code == 404


def test_export_tasks_csv(client, db_session):
    seed = _seed(db_session)
    client.post("/api/import/commit", files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")})

    resp = client.get("/api/export/tasks", params={"format": "csv"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    reader = csv.reader(io.StringIO(resp.content.decode("utf-8-sig")))
    rows = list(reader)
    assert rows[0][0] == "Employee"
    assert any(r[1] == "Design homepage" for r in rows[1:])


def test_export_tasks_xlsx(client, db_session):
    seed = _seed(db_session)
    client.post("/api/import/commit", files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")})

    resp = client.get("/api/export/tasks", params={"format": "xlsx"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb["Tasks"]["A1"].value == "Employee"


def test_export_tasks_invalid_format(client):
    resp = client.get("/api/export/tasks", params={"format": "pdf"})
    assert resp.status_code == 422


def test_export_tasks_filters_by_project(client, db_session):
    seed = _seed(db_session)
    other_project = Project(name="Other Project", status=ProjectStatus.ACTIVE)
    db_session.add(other_project)
    db_session.commit()
    db_session.add(Task(project_id=other_project.id, title="Other task"))
    db_session.add(Task(project_id=seed["project"].id, title="In scope task"))
    db_session.commit()

    resp = client.get("/api/export/tasks", params={"format": "csv", "project_id": seed["project"].id})
    text = resp.text
    assert "In scope task" in text
    assert "Other task" not in text


def test_export_tasks_unknown_project_404(client):
    resp = client.get("/api/export/tasks", params={"project_id": 999999})
    assert resp.status_code == 404


def test_export_dashboard_csv(client, db_session):
    seed = _seed(db_session)
    client.post("/api/import/commit", files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")})

    resp = client.get("/api/export/dashboard", params={"format": "csv"})
    assert resp.status_code == 200
    assert "Summary" in resp.text
    assert "Employee Performance" in resp.text


def test_export_dashboard_xlsx(client, db_session):
    seed = _seed(db_session)
    client.post("/api/import/commit", files={"file": ("tasks.csv", _csv_bytes(seed["project"].name), "text/csv")})

    resp = client.get("/api/export/dashboard", params={"format": "xlsx"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Summary" in wb.sheetnames


def test_export_dashboard_invalid_date_range(client):
    resp = client.get("/api/export/dashboard", params={"date_from": "2026-08-10", "date_to": "2026-08-01"})
    assert resp.status_code == 422
