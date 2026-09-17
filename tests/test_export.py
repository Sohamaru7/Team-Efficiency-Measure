"""Unit tests for app.services.import_export.export — CSV/XLSX generation for the task list and
the manager-dashboard report. Checks structural correctness (headers, row counts, section
titles) rather than byte-for-byte file content.
"""

import csv
import io
from datetime import date

from openpyxl import load_workbook

from app.models.enums import ProjectStatus, TaskPriority, TaskStatus, UserRole
from app.models.project import Project
from app.models.task import Task
from app.models.user import User
from app.services.analytics.dashboard import get_manager_dashboard
from app.services.import_export.export import (
    TASK_EXPORT_HEADERS,
    build_task_rows,
    export_dashboard_csv,
    export_dashboard_xlsx,
    export_tasks_csv,
    export_tasks_xlsx,
)


def _seed(db_session):
    manager = User(name="Manager Export", email="mgr-export@example.com", role=UserRole.MANAGER, department="Engineering")
    alice = User(name="Alice Export", email="alice-export@example.com", role=UserRole.EMPLOYEE, department="Engineering")
    db_session.add_all([manager, alice])
    db_session.commit()
    project = Project(name="Export Project", manager_id=manager.id, status=ProjectStatus.ACTIVE)
    db_session.add(project)
    db_session.commit()
    task = Task(
        project_id=project.id, assigned_to=alice.id, title="Design homepage",
        priority=TaskPriority.HIGH, status=TaskStatus.IN_PROGRESS, estimated_hours=5,
        start_date=date(2026, 8, 1), deadline=date(2026, 8, 10),
    )
    db_session.add(task)
    db_session.commit()
    return {"manager": manager, "alice": alice, "project": project, "task": task}


def test_build_task_rows_filters_by_project(db_session):
    seed = _seed(db_session)
    other_project = Project(name="Other Project", status=ProjectStatus.ACTIVE)
    db_session.add(other_project)
    db_session.commit()
    other_task = Task(project_id=other_project.id, title="Other task")
    db_session.add(other_task)
    db_session.commit()

    rows = build_task_rows(db_session, project_id=seed["project"].id)
    assert [t.id for t in rows] == [seed["task"].id]


def test_export_tasks_csv_round_trip_shape(db_session):
    seed = _seed(db_session)
    rows = build_task_rows(db_session)
    content = export_tasks_csv(rows)

    reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
    all_rows = list(reader)
    assert all_rows[0] == TASK_EXPORT_HEADERS
    assert all_rows[1][0] == "Alice Export"  # Employee
    assert all_rows[1][1] == "Design homepage"  # Task
    assert all_rows[1][2] == "Export Project"  # Project
    assert all_rows[1][3] == "high"
    assert all_rows[1][6] == "2026-08-01"  # Start Date
    assert all_rows[1][8] == "in_progress"  # Status


def test_export_tasks_xlsx_round_trip_shape(db_session):
    seed = _seed(db_session)
    rows = build_task_rows(db_session)
    content = export_tasks_xlsx(rows)

    wb = load_workbook(io.BytesIO(content))
    sheet = wb["Tasks"]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0] == tuple(TASK_EXPORT_HEADERS)
    assert values[1][0] == "Alice Export"
    assert values[1][1] == "Design homepage"


def test_export_tasks_unassigned_task_has_blank_employee(db_session):
    seed = _seed(db_session)
    unassigned = Task(project_id=seed["project"].id, title="No owner")
    db_session.add(unassigned)
    db_session.commit()

    rows = build_task_rows(db_session)
    content = export_tasks_csv(rows)
    reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
    all_rows = list(reader)
    unassigned_row = next(r for r in all_rows if r[1] == "No owner")
    assert unassigned_row[0] == ""


def test_export_dashboard_csv_has_all_sections(db_session):
    _seed(db_session)
    result = get_manager_dashboard(db_session)
    content = export_dashboard_csv(result)
    text = content.decode("utf-8-sig")
    for section in ["Summary", "Employee Performance", "Workload Distribution", "Project Progress", "Delayed Tasks", "Upcoming Deadlines"]:
        assert section in text


def test_export_dashboard_xlsx_has_one_sheet_per_section(db_session):
    _seed(db_session)
    result = get_manager_dashboard(db_session)
    content = export_dashboard_xlsx(result)
    wb = load_workbook(io.BytesIO(content))
    assert set(wb.sheetnames) == {
        "Summary", "Employee Performance", "Workload Distribution", "Project Progress", "Delayed Tasks", "Upcoming Deadlines",
    }
    summary_sheet = wb["Summary"]
    values = list(summary_sheet.iter_rows(values_only=True))
    assert values[0] == ("Metric", "Value")
