"""CSV/Excel export for manager reports (Phase 8). Two kinds of export, both reachable as
`?format=csv` or `?format=xlsx`:

- **Task export** (`build_task_rows` + `export_tasks_*`) — the raw task list in the same
  canonical column layout the importer understands, so an exported file can be edited and
  re-imported (a practical round trip, and a good way to sanity-check the import column
  mapping against real data).
- **Dashboard export** (`export_dashboard_*`) — the Manager Dashboard's own computed sections
  (summary KPIs, employee performance, workload, project progress, delayed tasks, upcoming
  deadlines), reusing `get_manager_dashboard` so an exported report always matches what the
  dashboard page shows for the same filters. CSV writes one section after another separated by
  a blank line; Excel writes one worksheet per section.
"""

from __future__ import annotations

import csv
import io
from datetime import date as date_type
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.enums import TaskStatus
from app.models.task import Task
from app.models.user import User
from app.services.analytics.dashboard_results import ManagerDashboardResult

TASK_EXPORT_HEADERS = [
    "Employee",
    "Task",
    "Project",
    "Priority",
    "Estimated Hours",
    "Actual Hours",
    "Start Date",
    "Deadline",
    "Status",
    "Quality",
    "Delay Reason",
]


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date_type):
        return value.isoformat()
    return str(value)


def _task_export_row(task: Task) -> list[str]:
    return [
        task.assignee.name if task.assignee else "",
        task.title,
        task.project.name if task.project else "",
        task.priority.value,
        _fmt(task.estimated_hours),
        _fmt(task.actual_hours),
        _fmt(task.start_date),
        _fmt(task.deadline),
        task.status.value,
        _fmt(task.quality_score),
        task.delay_reason or "",
    ]


def build_task_rows(
    db: Session,
    *,
    project_id: int | None = None,
    employee_id: int | None = None,
    department: str | None = None,
    status: TaskStatus | None = None,
) -> list[Task]:
    stmt = select(Task).options(selectinload(Task.project), selectinload(Task.assignee))
    if project_id is not None:
        stmt = stmt.where(Task.project_id == project_id)
    if employee_id is not None:
        stmt = stmt.where(Task.assigned_to == employee_id)
    if status is not None:
        stmt = stmt.where(Task.status == status)
    if department is not None:
        stmt = stmt.join(Task.assignee).where(User.department == department)
    stmt = stmt.order_by(Task.id)
    return list(db.scalars(stmt))


def export_tasks_csv(tasks: list[Task]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TASK_EXPORT_HEADERS)
    for task in tasks:
        writer.writerow(_task_export_row(task))
    return buf.getvalue().encode("utf-8-sig")


def export_tasks_xlsx(tasks: list[Task]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"
    sheet.append(TASK_EXPORT_HEADERS)
    for task in tasks:
        sheet.append(_task_export_row(task))
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _dashboard_sections(result: ManagerDashboardResult) -> list[tuple[str, list[str], list[list[str]]]]:
    summary = result.summary
    summary_rows = [
        ["As of", _fmt(summary.as_of)],
        ["Team size", _fmt(summary.team_size)],
        ["Overall team efficiency", _fmt(summary.overall_team_efficiency)],
        ["Tasks completed", _fmt(summary.tasks_completed)],
        ["Tasks overdue", _fmt(summary.tasks_overdue)],
        ["Tasks at risk", _fmt(summary.tasks_at_risk)],
        ["Average completion days", _fmt(summary.average_completion_days)],
        ["On-time completion rate", _fmt(summary.on_time_completion_rate)],
        ["Team workload utilization %", _fmt(summary.team_workload.utilization_pct)],
    ]

    employee_rows = [
        [row.name, row.department or "", _fmt(row.score.overall_score), _fmt(row.score.workload.utilization_pct)]
        for row in result.employee_performance
    ]

    workload_rows = [
        [row.name, row.department or "", _fmt(row.workload.utilization_pct), _fmt(row.workload.active_hours), _fmt(row.workload.capacity_hours)]
        for row in result.workload_distribution
    ]

    project_rows = [
        [row.name, _fmt(row.progress.progress_pct), _fmt(row.progress.completed_tasks), _fmt(row.progress.total_tasks)]
        for row in result.project_progress
    ]

    delayed_rows = [
        [row.title, row.project_name, row.assignee_name or "", row.status, _fmt(row.deadline), _fmt(row.delay_days), row.delay_reason or ""]
        for row in result.delayed_tasks
    ]

    upcoming_rows = [
        [row.title, row.project_name, row.assignee_name or "", row.status, row.priority, _fmt(row.deadline), _fmt(row.days_left)]
        for row in result.upcoming_deadlines
    ]

    return [
        ("Summary", ["Metric", "Value"], summary_rows),
        ("Employee Performance", ["Employee", "Department", "Overall Score", "Workload %"], employee_rows),
        ("Workload Distribution", ["Employee", "Department", "Utilization %", "Active Hours", "Capacity Hours"], workload_rows),
        ("Project Progress", ["Project", "Progress %", "Completed Tasks", "Total Tasks"], project_rows),
        ("Delayed Tasks", ["Task", "Project", "Assignee", "Status", "Deadline", "Delay Days", "Delay Reason"], delayed_rows),
        ("Upcoming Deadlines", ["Task", "Project", "Assignee", "Status", "Priority", "Deadline", "Days Left"], upcoming_rows),
    ]


def export_dashboard_csv(result: ManagerDashboardResult) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    for title, headers, rows in _dashboard_sections(result):
        writer.writerow([title])
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
        writer.writerow([])
    return buf.getvalue().encode("utf-8-sig")


def export_dashboard_xlsx(result: ManagerDashboardResult) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in _dashboard_sections(result):
        sheet = workbook.create_sheet(title[:31])  # Excel sheet-name limit
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()
