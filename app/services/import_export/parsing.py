"""Turns raw uploaded bytes into a plain `list[dict[str, str]]` of rows keyed by their original
file headers. This is the file-format boundary: CSV and Excel are handled here and nowhere
else, so everything downstream (column mapping, validation, duplicate detection, insertion)
works against the same plain-Python row shape regardless of source format.

This is also the intended extension point for a future Google Sheets integration (see the
Phase 8 README section): a `parse_google_sheet(...) -> ParsedFile` function returning the same
`ParsedFile` shape would plug into `importer.run_import` with no changes to any other module in
this package — the rest of the pipeline has no notion of "file" at all, only "headers + rows".
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.services.import_export.constants import MAX_IMPORT_ROWS
from app.services.import_export.errors import EmptyFileError, TooManyRowsError, UnsupportedFileTypeError


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, str]]


def parse_upload(filename: str, content: bytes) -> ParsedFile:
    """Dispatch on file extension. Raises an `ImportFileError` subclass for anything that isn't
    a readable CSV/XLSX with a header row and at least one data row.
    """
    name = (filename or "").lower()
    if name.endswith(".csv"):
        parsed = _parse_csv(content)
    elif name.endswith(".xlsx"):
        parsed = _parse_excel(content)
    elif name.endswith(".xls"):
        raise UnsupportedFileTypeError("Legacy .xls files are not supported — save as .xlsx or .csv.")
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {filename!r}. Use .csv or .xlsx.")

    if len(parsed.rows) > MAX_IMPORT_ROWS:
        raise TooManyRowsError(f"File has {len(parsed.rows)} data rows; the limit per import is {MAX_IMPORT_ROWS}.")
    return parsed


def _parse_csv(content: bytes) -> ParsedFile:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UnsupportedFileTypeError("Could not decode file as UTF-8 CSV.") from exc

    reader = csv.reader(io.StringIO(text))
    all_rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not all_rows:
        raise EmptyFileError("File has no header row.")

    headers = [h.strip() for h in all_rows[0]]
    data_rows = all_rows[1:]
    if not data_rows:
        raise EmptyFileError("File has a header row but no data rows.")

    rows = [_zip_row(headers, row) for row in data_rows]
    return ParsedFile(headers=headers, rows=rows)


def _parse_excel(content: bytes) -> ParsedFile:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several distinct exception types for a corrupt file
        raise UnsupportedFileTypeError("Could not read file as an Excel (.xlsx) workbook.") from exc

    sheet = workbook.active
    all_rows = [
        [_cell_to_str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    if not all_rows:
        raise EmptyFileError("File has no header row.")

    headers = [h.strip() for h in all_rows[0]]
    data_rows = all_rows[1:]
    if not data_rows:
        raise EmptyFileError("File has a header row but no data rows.")

    rows = [_zip_row(headers, row) for row in data_rows]
    return ParsedFile(headers=headers, rows=rows)


def _zip_row(headers: list[str], row: list[str]) -> dict[str, str]:
    # A short/ragged data row (fewer cells than headers) pads with "" rather than erroring —
    # a genuinely required field being blank is still caught by validation.py.
    padded = list(row) + [""] * (len(headers) - len(row))
    return {header: (padded[i] or "").strip() for i, header in enumerate(headers)}


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        # openpyxl returns every Excel date cell as a datetime (even a pure date, at midnight)
        # -- drop the zero time-of-day so a plain date cell round-trips as "2026-08-01", not
        # "2026-08-01T00:00:00".
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
